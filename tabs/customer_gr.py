"""Tab C: Đổ giá nhóm KH"""
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import threading
import os
from types import SimpleNamespace

from helper.functions import (
    BG_COLOR, TEXT_COLOR, CARD_BG, STATUS_BG,
    ROW_EVEN, ROW_ODD, BTN_SUCCESS, BTN_SUCCESS_HOVER,
    CUSTOMER_GROUP_MAP,
    _get_or_create_sap_session, reset_sap_session,
    on_enter_btn, on_leave_btn,
    bind_tree_shortcuts,
)


def build(root, frame_c):
    """Tạo toàn bộ UI và logic cho Tab C. Trả về namespace chứa widgets cần thiết."""

    CURRENT_DATA_C = []

    # ---- Loading overlay ----
    _state = {"frame": None, "label": None, "anim_id": None, "dots": 0}

    def show_loading_c(msg="Đang tải dữ liệu"):
        hide_loading_c()
        _state["dots"] = 0
        _state["frame"] = tk.Frame(frame_c, bg=BG_COLOR)
        _state["frame"].place(relx=0.5, rely=0.55, anchor=tk.CENTER, width=300, height=80)
        _state["label"] = tk.Label(
            _state["frame"], text=msg, font=("Segoe UI", 12), bg=BG_COLOR, fg=TEXT_COLOR)
        _state["label"].pack(expand=True)
        _animate(msg)

    def _animate(base_msg):
        if _state["label"] and _state["label"].winfo_exists():
            _state["dots"] = (_state["dots"] % 3) + 1
            _state["label"].config(text=base_msg + "." * _state["dots"])
            _state["anim_id"] = root.after(400, lambda: _animate(base_msg))

    def hide_loading_c():
        if _state["anim_id"]:
            root.after_cancel(_state["anim_id"])
            _state["anim_id"] = None
        if _state["frame"]:
            _state["frame"].destroy()
            _state["frame"] = None
            _state["label"] = None

    # ---- Logic functions ----
    def download_template_c():
        """Tạo và lưu file Excel template mẫu."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            messagebox.showerror(
                "Thiếu thư viện",
                "Cần cài đặt thư viện openpyxl.\n\nChạy lệnh: pip install openpyxl")
            return
        try:
            save_path = filedialog.asksaveasfilename(
                parent=root, title="Lưu file template",
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile="template_do_gia_nhom_kh.xlsx",
                initialdir=os.path.join(os.path.expanduser("~"), "Desktop"),
            )
            if not save_path:
                return
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "DuLieu"
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
            header_align = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"),
            )
            headers = ["Mã SP", "Giá"]
            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border
            sample_data = [("100001", 50000), ("100002", 75000), ("100003", 120000)]
            data_font = Font(name="Segoe UI", size=11)
            for row_idx, (masp, gia) in enumerate(sample_data, start=2):
                c1 = ws.cell(row=row_idx, column=1, value=masp)
                c1.font = data_font
                c1.border = thin_border
                c1.alignment = Alignment(horizontal="center")
                c2 = ws.cell(row=row_idx, column=2, value=gia)
                c2.font = data_font
                c2.border = thin_border
                c2.number_format = "#,##0"
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 20
            wb.save(save_path)
            lbl_status_c.config(text=f"  Đã lưu template: {os.path.basename(save_path)}")
            messagebox.showinfo("Thành công", f"Đã lưu file template tại:\n{save_path}")
        except Exception as ex:
            messagebox.showerror("Lỗi", f"Không thể tạo file template:\n{ex}")

    def load_excel_file_c():
        """Mở file Excel và đọc dữ liệu mã SP + giá vào bảng."""
        try:
            import openpyxl  # noqa: F811
        except ImportError:
            messagebox.showerror(
                "Thiếu thư viện",
                "Cần cài đặt thư viện openpyxl.\n\nChạy lệnh: pip install openpyxl")
            return
        file_path = filedialog.askopenfilename(
            parent=root, title="Chọn file Excel",
            filetypes=[("Excel files", "*.xlsx *.xls")],
        )
        if not file_path:
            return
        show_loading_c("Đang đọc file Excel")
        root.update_idletasks()

        def _read():
            try:
                import openpyxl  # noqa: F811
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws = wb.active
                data = []
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if row is None or len(row) < 2:
                        continue
                    masp_raw, gia_raw = row[0], row[1]
                    if masp_raw is None or gia_raw is None:
                        continue
                    masp = str(masp_raw).strip()
                    if not masp:
                        continue
                    try:
                        gia = float(str(gia_raw).replace(",", "").strip())
                    except (ValueError, TypeError):
                        continue
                    data.append({"masp": masp, "gia": gia})
                wb.close()
                CURRENT_DATA_C.clear()
                CURRENT_DATA_C.extend(data)
                root.after(0, lambda: hide_loading_c())
                root.after(0, lambda: _show_results_c(data, file_path))
            except Exception as ex:
                CURRENT_DATA_C.clear()
                root.after(0, lambda: hide_loading_c())
                root.after(0, lambda: lbl_status_c.config(text=f"  Lỗi đọc file: {ex}"))
                root.after(0, lambda: messagebox.showerror("Lỗi", f"Không thể đọc file Excel:\n{ex}"))
        threading.Thread(target=_read, daemon=True).start()

    def _show_results_c(data, file_path=""):
        for row in tree_c.get_children():
            tree_c.delete(row)
        for i, item in enumerate(data, start=1):
            gia_fmt = f"{item['gia']:,.0f}"
            tag = "even" if i % 2 == 0 else "odd"
            tree_c.insert("", tk.END, values=(i, item["masp"], gia_fmt), tags=(tag,))
        fname = os.path.basename(file_path) if file_path else ""
        lbl_status_c.config(text=f"  Đã tải {len(data)} sản phẩm từ {fname}")

    def do_update_price_c():
        selected = tree_c.selection()
        if not selected:
            messagebox.showwarning("Chưa chọn dòng", "Vui lòng chọn ít nhất một dòng trong bảng để đổ giá.")
            return
        group_label = cmb_group.get()
        group_code = CUSTOMER_GROUP_MAP.get(group_label, "")
        if not group_code:
            messagebox.showwarning("Chưa chọn nhóm KH", "Vui lòng chọn nhóm khách hàng.")
            return
        items = []
        for sel in selected:
            values = tree_c.item(sel, "values")
            items.append({
                "masp": values[1],
                "gia": values[2].replace(",", "").strip(),
                "gia_fmt": values[2],
            })
        count = len(items)
        MAX_SHOW = 10
        detail_lines = [f"  {i+1}. SP: {it['masp']}  |  Giá: {it['gia_fmt']}" for i, it in enumerate(items[:MAX_SHOW])]
        if count > MAX_SHOW:
            detail_lines.append(f"  ... và {count - MAX_SHOW} dòng nữa")
        detail = "\n".join(detail_lines)
        confirm = messagebox.askyesno(
            "Xác nhận đổ giá",
            f"Nhóm KH: {group_label}\n\nĐổ giá {count} sản phẩm?\n\n{detail}")
        if not confirm:
            return
        btn_update_c.config(state=tk.DISABLED, text="Đang xử lý...")
        lbl_status_c.config(text=f"  Đang xử lý {count} dòng...")
        root.update_idletasks()

        def _run():
            try:
                from helper.vk11 import updatematerialgr
                session = _get_or_create_sap_session()
                if session == "MAX_SESSIONS":
                    root.after(0, lambda: messagebox.showwarning(
                        "Quá giới hạn session",
                        "Số session SAP đã đạt giới hạn tối đa.\n\nVui lòng đóng bớt các cửa sổ SAP đang mở rồi thử lại."))
                    return
                if session is None:
                    root.after(0, lambda: messagebox.showerror("Lỗi", "Không thể kết nối SAP. Kiểm tra lại cấu hình."))
                    return
                mat_list = [{"masp": it["masp"], "gia": it["gia"]} for it in items]
                results = updatematerialgr(session, group_code, mat_list)
                success = 0
                errors = []
                for i, r in enumerate(results):
                    if r["status"] == "ok":
                        success += 1
                    else:
                        errors.append(f"{r['masp']}: {r['msg']}")
                    root.after(0, lambda p=i+1: lbl_status_c.config(text=f"  Đã xử lý {p}/{count}..."))
                msg = f"Thành công: {success}/{count}"
                if errors:
                    unique_msgs = set()
                    for e in errors:
                        if "| SAP: " in e:
                            unique_msgs.add(e.split("| SAP: ")[-1])
                        else:
                            unique_msgs.add(e.split(": ", 1)[-1] if ": " in e else e)
                    if len(unique_msgs) == 1:
                        msg += f"\n\nLỗi ({len(errors)} SP): {unique_msgs.pop()}"
                    else:
                        msg += f"\n\nLỗi ({len(errors)}):\n" + "\n".join(errors)
                    root.after(0, lambda m=msg: messagebox.showwarning("Kết quả", m))
                root.after(0, lambda: lbl_status_c.config(text=f"  Hoàn tất: {success}/{count} thành công"))
            except Exception as ex:
                reset_sap_session()
                root.after(0, lambda e=ex: messagebox.showerror("Lỗi", f"Đổ giá thất bại:\n{e}"))
                root.after(0, lambda e=ex: lbl_status_c.config(text=f"  Lỗi: {e}"))
            finally:
                root.after(0, lambda: btn_update_c.config(state=tk.NORMAL, text="Đổ giá SAP"))
        threading.Thread(target=_run, daemon=True).start()

    # ---- UI: Card tìm kiếm ----
    card_search_c = tk.Frame(frame_c, bg=CARD_BG, bd=0, highlightthickness=1, highlightbackground="#dce6f0")
    card_search_c.pack(fill=tk.X, padx=16, pady=(16, 8))

    inner_search_c = tk.Frame(card_search_c, bg=CARD_BG)
    inner_search_c.pack(fill=tk.X, padx=20, pady=16)

    tk.Label(inner_search_c, text="Nhóm khách hàng", font=("Segoe UI", 10, "bold"), bg=CARD_BG, fg=TEXT_COLOR).grid(
        row=0, column=0, padx=(0, 24), sticky=tk.W)

    cmb_group = ttk.Combobox(
        inner_search_c, values=list(CUSTOMER_GROUP_MAP.keys()), width=30, font=("Segoe UI", 11),
        state="readonly")
    cmb_group.current(0)
    cmb_group.grid(row=1, column=0, padx=(0, 16), pady=(4, 0), ipady=3, sticky=tk.W)

    BTN_IMPORT = "#2980b9"
    BTN_IMPORT_HOVER = "#3498db"
    BTN_TEMPLATE = "#8e44ad"
    BTN_TEMPLATE_HOVER = "#9b59b6"

    btn_load_excel = tk.Button(
        inner_search_c, text="📂 Nạp file Excel", font=("Segoe UI", 10, "bold"),
        bg=BTN_IMPORT, fg="white", activebackground=BTN_IMPORT_HOVER, activeforeground="white",
        relief=tk.FLAT, cursor="hand2", padx=12, pady=4, command=load_excel_file_c)
    btn_load_excel.grid(row=1, column=1, padx=(0, 8), pady=(4, 0), ipady=3)
    btn_load_excel.bind("<Enter>", lambda e: on_enter_btn(e, BTN_IMPORT_HOVER))
    btn_load_excel.bind("<Leave>", lambda e: on_leave_btn(e, BTN_IMPORT))

    btn_template = tk.Button(
        inner_search_c, text="📥 Tải template mẫu", font=("Segoe UI", 10, "bold"),
        bg=BTN_TEMPLATE, fg="white", activebackground=BTN_TEMPLATE_HOVER, activeforeground="white",
        relief=tk.FLAT, cursor="hand2", padx=12, pady=4, command=download_template_c)
    btn_template.grid(row=1, column=2, padx=(0, 8), pady=(4, 0), ipady=3)
    btn_template.bind("<Enter>", lambda e: on_enter_btn(e, BTN_TEMPLATE_HOVER))
    btn_template.bind("<Leave>", lambda e: on_leave_btn(e, BTN_TEMPLATE))

    btn_update_c = tk.Button(
        inner_search_c, text="Đổ giá SAP", font=("Segoe UI", 10, "bold"),
        bg=BTN_SUCCESS, fg="white", activebackground=BTN_SUCCESS_HOVER, activeforeground="white",
        relief=tk.FLAT, cursor="hand2", padx=16, pady=4, command=do_update_price_c)
    btn_update_c.grid(row=1, column=3, padx=(0, 8), pady=(4, 0), ipady=3)
    btn_update_c.bind("<Enter>", lambda e: on_enter_btn(e, BTN_SUCCESS_HOVER))
    btn_update_c.bind("<Leave>", lambda e: on_leave_btn(e, BTN_SUCCESS))

    # ---- UI: Card bảng dữ liệu ----
    card_table_c = tk.Frame(frame_c, bg=CARD_BG, bd=0, highlightthickness=1, highlightbackground="#dce6f0")
    card_table_c.pack(fill=tk.BOTH, expand=True, padx=16, pady=(8, 8))

    tbl_frame_c = tk.Frame(card_table_c, bg=CARD_BG)
    tbl_frame_c.pack(fill=tk.BOTH, expand=True, padx=0, pady=0)

    columns_c = ("stt", "masp", "gia")
    tree_c = ttk.Treeview(tbl_frame_c, columns=columns_c, show="headings", height=14, style="Custom.Treeview", selectmode="extended")

    tree_c.heading("stt", text="STT")
    tree_c.heading("masp", text="Mã SP")
    tree_c.heading("gia", text="Giá (VNĐ)")

    tree_c.column("stt", width=80, anchor=tk.CENTER, minwidth=50)
    tree_c.column("masp", width=250, anchor=tk.CENTER, minwidth=120)
    tree_c.column("gia", width=250, anchor=tk.E, minwidth=120)

    tree_c.tag_configure("odd", background=ROW_ODD)
    tree_c.tag_configure("even", background=ROW_EVEN)

    tbl_frame_c.columnconfigure(0, weight=1)
    tbl_frame_c.rowconfigure(0, weight=1)
    scrollbar_c = ttk.Scrollbar(tbl_frame_c, orient=tk.VERTICAL, command=tree_c.yview, style="Custom.Vertical.TScrollbar")
    tree_c.configure(yscrollcommand=scrollbar_c.set)
    tree_c.grid(row=0, column=0, sticky="nsew")
    scrollbar_c.grid(row=0, column=1, sticky="ns")

    # ---- Status bar ----
    status_bar_c = tk.Frame(frame_c, bg=STATUS_BG, height=32)
    status_bar_c.pack(fill=tk.X, padx=16, pady=(0, 12))
    status_bar_c.pack_propagate(False)

    lbl_status_c = tk.Label(status_bar_c, text="", font=("Segoe UI", 9), bg=STATUS_BG, fg="#5d6d7e", anchor=tk.W)
    lbl_status_c.pack(fill=tk.X, padx=8, pady=5)

    # ---- Bindings ----
    paste_and_select_c = bind_tree_shortcuts(root, tree_c, masp_col_index=1, lbl_status=lbl_status_c)

    def tree_enter_c(event):
        if tree_c.selection():
            do_update_price_c()
        return "break"
    tree_c.bind("<Return>", tree_enter_c)

    # ---- Return namespace ----
    return SimpleNamespace(
        tree_c=tree_c,
        lbl_status_c=lbl_status_c,
        paste_and_select_c=paste_and_select_c,
    )
